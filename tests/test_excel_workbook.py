"""The excel template's formula values and agent tools (excel/workbook.py),
and how reader.py hands them to the editor.

workbook.py and reader.py are runPython targets, not package modules —
loaded via importlib and driven directly against tmp_path files, like
test_excel_readonly.py does.

Under test:
* Formula values: Excel's cached value when there is one (pycel never runs);
  otherwise pycel's, for every formula cell of the workbook in ONE bounded
  worker run, cached on disk per file version. An error comes through like
  Excel's; a formula pycel can't evaluate is absent (the grid keeps its own
  engine's answer) and counted as `unevaluated` by the tools; a worker that
  hangs is cut off by its timeout.
* Dates: DATE/TODAY/… and a `ref ± N` offset off a date are dates; a day
  count (`TODAY()-A2`) stays a number.
* The editor: `load` sends each sheet's formula values as `computed`
  ({"r,c": value}, 0-based); `compute` evaluates the UNSAVED sheets it is sent.
* The agent tools — sheets, describe, query, cells — over the computed
  workbook; query's SQL is sandboxed (no files, no network) and time-bounded.
"""
import importlib.util
import math
import os
import time

import pytest

# Not every test environment installs the `bundled` extra (test_excel_readonly.py
# notes the same gap) — skip the file rather than fail collection.
openpyxl = pytest.importorskip("openpyxl")
pytest.importorskip("pycel")


def _load(name):
    path = os.path.join("fused_render", "templates", "excel", name)
    spec = importlib.util.spec_from_file_location(name[:-3], path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


@pytest.fixture(autouse=True)
def _isolated_cache(tmp_path, monkeypatch):
    # The formula cache lives under ~ by design; every test starts cold.
    monkeypatch.setenv("HOME", str(tmp_path / "home"))


def _write(path, rows, sheet_name="Sheet1"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    for row in rows:
        ws.append(row)
    wb.save(path)


def _values(path, sheet="Sheet1"):
    return _load("workbook.py").evaluate(str(path))[sheet]


def _no_spawn(mod, monkeypatch):
    monkeypatch.setattr(mod.subprocess, "run",
                        lambda *a, **k: pytest.fail("the workbook was evaluated again"))


# ---------------------------------------------------------------------------
# formula values
# ---------------------------------------------------------------------------

def test_formula_with_no_cached_value_is_evaluated(tmp_path):
    path = tmp_path / "book.xlsx"
    # openpyxl never computes formulas on save — the shape of a workbook no
    # Excel ever opened (built by a script, or saved by this editor).
    _write(path, [["hours", "amount"], [3, "=A2*50"], [0, "=1/A3"]])

    assert _values(path) == {(1, 1): 150, (2, 1): "#DIV/0!"}


def test_cross_sheet_lookups_are_evaluated(tmp_path):
    """The shapes the editor's in-browser engine can't compute."""
    path = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    rates = wb.active
    rates.title = "Rate table"
    rates.append(["tier", "rate"])
    rates.append(["base", 0.1])
    rates.append(["high", 0.25])
    calc = wb.create_sheet("Calc")
    calc.append(["tier", "rate"])
    calc.append(["high", "=INDEX('Rate table'!B2:B3,MATCH(A2,'Rate table'!A2:A3,0))"])
    wb.save(path)

    assert _values(path, "Calc") == {(1, 1): 0.25}


def test_cached_value_is_used_without_running_pycel(tmp_path, monkeypatch):
    path = tmp_path / "book.xlsx"
    _write(path, [["a", "b"], [1, "=A2*10"]])
    # openpyxl's writer never emits a cached <v> for a formula cell, so one is
    # spliced into the sheet XML by hand — the shape a real Excel save produces.
    import zipfile

    with zipfile.ZipFile(path) as zin:
        contents = {n: zin.read(n) for n in zin.namelist()}
    sheet_xml = contents["xl/worksheets/sheet1.xml"].decode()
    contents["xl/worksheets/sheet1.xml"] = sheet_xml.replace(
        "<f>A2*10</f>", "<f>A2*10</f><v>10</v>", 1).encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, data in contents.items():
            zout.writestr(n, data)

    mod = _load("workbook.py")
    _no_spawn(mod, monkeypatch)
    assert mod.evaluate(str(path))["Sheet1"] == {(1, 1): 10}


def test_unsupported_formula_is_absent_not_a_crash(tmp_path):
    path = tmp_path / "book.xlsx"
    _write(path, [["a", "b", "c"], [1, "=NOT_A_REAL_FUNCTION(A2)", "=A2+1"]])

    assert _values(path) == {(1, 2): 2}
    sheet = _load("workbook.py").sheets(str(path))["sheets"][0]
    assert sheet["formulas"] == 2 and sheet["unevaluated"] == 1


def test_sheet_name_with_spaces_and_quotes(tmp_path):
    path = tmp_path / "book.xlsx"
    _write(path, [["a", "b"], [3, "=A2*2"]], sheet_name="Bob's invoice")

    assert _values(path, "Bob's invoice") == {(1, 1): 6}


def test_date_formulas_are_dates_and_day_counts_are_numbers(tmp_path):
    import datetime

    path = tmp_path / "book.xlsx"
    _write(path, [
        ["issued", "due", "back", "elapsed", "span"],
        ["=DATE(2024,3,5)", "=A2+30", "=$A$2-10", "=DATE(2024,3,15)-A2",
         "=DATE(2024,3,5)-DATE(2024,1,1)"],
    ])

    assert _values(path) == {
        (1, 0): datetime.datetime(2024, 3, 5),
        (1, 1): datetime.datetime(2024, 4, 4),   # an offset off a date is a date
        (1, 2): datetime.datetime(2024, 2, 24),  # through an absolute $A$2 too
        (1, 3): 10,                              # a day count stays a number
        (1, 4): 64,
    }


def test_the_workbook_is_evaluated_once_per_version(tmp_path, monkeypatch):
    path = tmp_path / "book.xlsx"
    _write(path, [["a", "b"], [2, "=A2*10"]])
    assert _values(path) == {(1, 1): 20}

    mod = _load("workbook.py")  # a fresh process reads the disk cache instead
    _no_spawn(mod, monkeypatch)
    assert mod.evaluate(str(path))["Sheet1"] == {(1, 1): 20}

    monkeypatch.undo()
    monkeypatch.setenv("HOME", str(tmp_path / "home"))
    _write(path, [["a", "b"], [3, "=A2*10"]])
    os.utime(path, ns=(time.time_ns(), time.time_ns() + 10**9))  # a distinct mtime, however coarse the fs
    assert _load("workbook.py").evaluate(str(path))["Sheet1"] == {(1, 1): 30}


def test_a_scratch_file_is_not_cached(tmp_path):
    path = tmp_path / "book.xlsx"
    _write(path, [["a", "b"], [2, "=A2*10"]])

    _load("workbook.py").evaluate(str(path), cache=False)
    assert not (tmp_path / "home").exists()


def test_a_hung_worker_times_out(tmp_path):
    path = tmp_path / "book.xlsx"
    _write(path, [["a", "b"], [1, "=A2*10"]])
    slow = tmp_path / "slow_worker.py"
    slow.write_text("import sys, time\nsys.stdin.read()\ntime.sleep(60)\n")

    mod = _load("workbook.py")
    mod._PYCEL_WORKER = str(slow)
    mod._PYCEL_TIMEOUT = 0.5
    started = time.monotonic()
    assert mod.evaluate(str(path))["Sheet1"] == {}
    assert time.monotonic() - started < 5


def test_non_finite_results_are_excel_num_errors_not_invalid_json():
    worker = _load("_pycel_worker.py")
    assert worker._plain(math.inf) == "#NUM!"
    assert worker._plain(-math.inf) == "#NUM!"
    assert worker._plain(math.nan) is None
    assert _load("workbook.py")._jsonify(math.inf) == "#NUM!"


# ---------------------------------------------------------------------------
# the editor: load + compute
# ---------------------------------------------------------------------------

def test_load_sends_each_sheets_formula_values(tmp_path):
    path = tmp_path / "book.xlsx"
    _write(path, [["hours", "amount"], [3, "=A2*50"], [2, "=A3*50"]])

    sheet = _load("reader.py").main(action="load", file=str(path))["sheets"][0]
    assert sheet["rows"][1] == [3, "=A2*50"]  # the editor still gets the formula
    assert sheet["computed"] == {"1,1": 150, "2,1": 100}


def test_compute_evaluates_the_unsaved_sheets(tmp_path):
    path = tmp_path / "book.xlsx"
    _write(path, [["x"], [1]])
    before = os.path.getmtime(path)
    unsaved = [
        {"name": "Inputs", "rows": [["rate"], ["0.5"]]},
        {"name": "Model", "rows": [["value"], ["=Inputs!A2*10"], ["=HYPERLINK(\"#Inputs!A1\",\"go\")"]]},
    ]

    out = _load("reader.py").main(action="compute", file=str(path), data=__import__("json").dumps(unsaved))
    assert out["computed"] == [{}, {"1,0": 5}]  # HYPERLINK is pycel-less: the grid's own engine shows it
    assert os.path.getmtime(path) == before     # nothing was written to the workbook


# ---------------------------------------------------------------------------
# agent tools
# ---------------------------------------------------------------------------

def _sales(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Q1 sales"
    ws.append(["region", "units", "price", "revenue"])
    for i, (region, units, price) in enumerate(
            [("N", 3, 10.5), ("S", 5, 2.5), ("N", 1, 4), ("E", 0, 1)], start=2):
        ws.append([region, units, price, f"=B{i}*C{i}"])
    ws.append(["total", "=SUM(B2:B5)", None, "=1/0"])
    other = wb.create_sheet("Notes")
    other.append(["note", "note"])  # a repeated header stays addressable
    other.append(["hello", "again"])
    wb.save(path)


def test_sheets_lists_tables_with_types(tmp_path):
    path = tmp_path / "sales.xlsx"
    _sales(path)

    q1, notes = _load("workbook.py").sheets(str(path))["sheets"]
    assert q1["name"] == "Q1 sales" and q1["rows"] == 5 and q1["unevaluated"] == 0
    assert {c["name"]: c["type"] for c in q1["columns"]} == {
        "region": "string", "units": "int64", "price": "double", "revenue": "double"}
    assert [c["name"] for c in notes["columns"]] == ["note", "note_2"]


def test_describe_gives_column_statistics(tmp_path):
    path = tmp_path / "sales.xlsx"
    _sales(path)

    out = _load("workbook.py").describe(str(path), columns="units,revenue,region")
    units, revenue, region = out["columns"]
    assert units == {"column": "units", "type": "int64", "count": 5, "nulls": 0, "distinct": 5,
                     "min": 0, "max": 9, "mean": 3.6, "median": 3.0,
                     "std": pytest.approx(3.577708), "sum": 18}
    # the #DIV/0! total is an error, not a number: NULL in SQL, and counted
    assert revenue["sum"] == 48.0 and revenue["nulls"] == 1 and revenue["errors"] == 1
    assert region["top"][0] == {"value": "N", "count": 2}

    with pytest.raises(ValueError, match="nope"):
        _load("workbook.py").describe(str(path), columns="nope")


def test_query_runs_sql_over_every_sheet(tmp_path):
    path = tmp_path / "sales.xlsx"
    _sales(path)

    out = _load("reader.py").main(action="query", file=str(path), sql=(
        'SELECT region, sum(revenue) AS revenue FROM "Q1 sales" '
        "WHERE region <> 'total' GROUP BY 1 ORDER BY 2 DESC"))
    assert out["columns"] == ["region", "revenue"]
    assert out["rows"] == [["N", 35.5], ["S", 12.5], ["E", 0.0]]
    assert out["truncated"] is False

    capped = _load("workbook.py").query(str(path), 'SELECT * FROM "Q1 sales"', limit=2)
    assert len(capped["rows"]) == 2 and capped["truncated"] is True


@pytest.mark.parametrize("sql", [
    "SELECT * FROM read_csv('/etc/hosts')",
    "COPY \"Q1 sales\" TO 'leak.csv'",
    "ATTACH 'other.db'",
    "SET enable_external_access = true",
])
def test_query_cannot_touch_files_or_reopen_the_sandbox(tmp_path, sql):
    import duckdb

    path = tmp_path / "sales.xlsx"
    _sales(path)
    with pytest.raises(duckdb.Error):
        _load("workbook.py").query(str(path), sql)


def test_query_is_time_bounded(tmp_path):
    path = tmp_path / "sales.xlsx"
    _sales(path)
    mod = _load("workbook.py")
    mod._QUERY_TIMEOUT = 0.3

    started = time.monotonic()
    with pytest.raises(TimeoutError):
        mod.query(str(path), "SELECT count(*) FROM range(100000) a, range(100000) b WHERE a.range + b.range = 7")
    assert time.monotonic() - started < 5


def test_cells_returns_values_and_formulas(tmp_path):
    path = tmp_path / "sales.xlsx"
    _sales(path)

    out = _load("reader.py").main(action="cells", file=str(path), ref="'Q1 sales'!B6:D6")
    assert out["sheet"] == "Q1 sales"
    assert out["cells"] == [[
        {"ref": "B6", "value": 9, "formula": "=SUM(B2:B5)"},
        {"ref": "C6", "value": None, "formula": None},
        {"ref": "D6", "value": "#DIV/0!", "formula": "=1/0"},
    ]]


def test_the_tools_refuse_an_oversized_workbook_by_name(tmp_path, monkeypatch):
    path = tmp_path / "sales.xlsx"
    _sales(path)
    mod = _load("workbook.py")
    monkeypatch.setattr(mod, "_SNAPSHOT_MAX_BYTES", 1)
    with pytest.raises(mod.TooLarge, match="MB"):
        mod.describe(str(path))


# ---------------------------------------------------------------------------
# review follow-ups
# ---------------------------------------------------------------------------

def _splice_cached(path, sheet_xml_name, after, cached):
    import zipfile

    with zipfile.ZipFile(path) as zin:
        contents = {n: zin.read(n) for n in zin.namelist()}
    xml = contents[sheet_xml_name].decode()
    assert after in xml, xml
    contents[sheet_xml_name] = xml.replace(after, after + f"<v>{cached}</v>", 1).encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, data in contents.items():
            zout.writestr(n, data)


def test_an_array_formula_is_a_formula_with_its_cached_value(tmp_path):
    from openpyxl.worksheet.formula import ArrayFormula

    path = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["a", "b", "total"])
    ws.append([1, 2, None])
    ws.append([3, 4, None])
    ws["C2"] = ArrayFormula("C2", "=SUM(A2:A3*B2:B3)")
    wb.save(path)
    _splice_cached(path, "xl/worksheets/sheet1.xml", "SUM(A2:A3*B2:B3)</f>", 14)

    assert _values(path) == {(1, 2): 14}
    sheet = _load("reader.py").main(action="load", file=str(path))["sheets"][0]
    assert sheet["rows"][1][2] == "=SUM(A2:A3*B2:B3)"  # the formula, not the object's repr
    assert sheet["computed"] == {"1,2": 14}


def test_an_offset_off_a_date_formatted_formula_is_a_date(tmp_path):
    import datetime

    path = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["issued", "due", "lookup"])
    ws.append(["=C2", "=A2+30", 45356])  # 45356 = 2024-03-05; A2 is a plain ref, formatted as a date
    ws["A2"].number_format = "yyyy-mm-dd"
    wb.save(path)

    values = _values(path)
    assert values[(1, 0)] == datetime.datetime(2024, 3, 5)
    assert values[(1, 1)] == datetime.datetime(2024, 4, 4)


def test_evaluate_reads_only_the_sheets_asked_for(tmp_path, monkeypatch):
    path = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    small = wb.active
    small.title = "Small"
    small.append(["n", "total"])
    small.append([2, "=SUM(Big!A1:A3)"])
    big = wb.create_sheet("Big")
    for i in range(3):
        big.append([i + 1])
    wb.save(path)

    mod = _load("workbook.py")
    assert mod.evaluate(str(path), sheets=["Small"]) == {"Small": {(1, 1): 6}}  # pycel still sees Big

    # The editor's load, with "Big" on the Parquet path, reads only "Small".
    reader = _load("reader.py")
    monkeypatch.setattr(reader, "_is_big", lambda nr, nc: nr == 3)
    asked = []
    real = reader._workbook().evaluate
    monkeypatch.setattr(reader._workbook(), "evaluate",
                        lambda file, **kw: asked.append(kw.get("sheets")) or real(file, **kw))
    sheets = reader.main(action="load", file=str(path))["sheets"]
    assert asked == [["Small"]]
    assert sheets[0]["computed"] == {"1,1": 6} and sheets[1]["big"] is True

    assert reader.main(action="formulas", file=str(path))["computed"] == [{"1,1": 6}, {}]


def test_a_workbook_of_only_big_sheets_evaluates_nothing(tmp_path, monkeypatch):
    path = tmp_path / "book.xlsx"
    _write(path, [["a", "b"], [1, "=A2*10"]])

    reader = _load("reader.py")
    monkeypatch.setattr(reader, "_is_big", lambda nr, nc: True)
    monkeypatch.setattr(reader._workbook(), "evaluate",
                        lambda *a, **k: pytest.fail("a big sheet was read into memory"))
    assert reader.main(action="load", file=str(path))["sheets"][0]["big"] is True
