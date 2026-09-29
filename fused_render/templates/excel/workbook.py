"""Formula values and agent tools for the excel template.

Two callers, one engine:

* **The editor** (reader.py's `load` / `compute`) asks `evaluate()` for the
  value of every formula cell, so the grid shows what Excel would even where
  its own in-browser engine can't compute a formula (INDEX/MATCH, cross-sheet
  references, …) — including in a workbook no Excel ever computed.
* **An agent** calls the tools below, published over MCP by `mcp.toml` beside
  this file (`fused app serve <this folder>`) and over /api/run through
  reader.py's actions of the same names:

    sheets(file)                   every sheet: name, row count, typed columns
    describe(file, sheet, columns) per-column statistics
    query(file, sql, limit)        read-only DuckDB SQL; every sheet is a table
    cells(file, ref, sheet)        an A1 range: value AND formula of each cell

Formula values. `data_only=True` gives the value Excel cached at its last
save. A workbook built by openpyxl — or saved by this editor, which writes
formulas without results — has none, so those cells are evaluated with pycel:
all of them in one run of `_pycel_worker.py`, a subprocess with its own
timeout, so a workbook pycel chokes on still opens (its formula cells just
stay uncomputed). The results are cached on disk per file version, so the next
open, or the next agent call, reads them instead. A cell pycel cannot evaluate
is counted in `unevaluated` — reported, never guessed.

Dates. pycel returns Excel day serials; a computed cell is a date when Excel
would show one: its own number format is a date format, the whole formula is
one date-function call (`DATE(…)`, `TODAY()`, …), or it is a plain `ref ± N`
offset off a date cell. `TODAY()-A2` is a day count and stays a number.

For the tools, the first row of a sheet is its header; blank header cells get
positional names (`col3`), and repeated ones a suffix (`Amount_2`), so every
column is addressable as a SQL identifier.
"""
import datetime
import decimal
import hashlib
import json
import math
import os
import re
import subprocess
import sys
import threading
from collections import OrderedDict

import openpyxl
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, range_boundaries
from openpyxl.utils.datetime import from_excel

_PYCEL_WORKER = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_pycel_worker.py")
# One run evaluates every uncached formula in the workbook, cached per file
# version — a one-time cost, bounded well inside the 60s a reader call gets.
_PYCEL_TIMEOUT = 20.0

# The tools hold the whole workbook in memory; past this size they refuse by
# name rather than thrash (the editor itself pages big sheets through Parquet).
_SNAPSHOT_MAX_BYTES = 25 * 1024 * 1024

# Formula results survive the process here, keyed by file version (see
# _formula_cache_path). ~ is expanded per call, never at import.
_CACHE_DIR = os.path.join("~", ".fused-render", "cache", "excel", "formulas")

# SQL is bounded: wall clock (interrupted), memory, threads, and no filesystem
# or network at all.
_QUERY_TIMEOUT = 10.0
_QUERY_MAX_ROWS = 10_000
_CELLS_MAX = 10_000
_TOP_VALUES = 5

_EXCEL_ERRORS = frozenset((
    "#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!",
    "#SPILL!", "#CALC!", "#GETTING_DATA",
))
# Whole-formula date calls (see module docstring, "Dates").
_DATE_FUNCS = ("DATE(", "TODAY(", "NOW(", "EDATE(", "EOMONTH(", "WORKDAY(", "WORKDAY.INTL(")
_OFFSET_RE = re.compile(
    r"^=?\s*(?:(\$?[A-Za-z]{1,3}\$?[0-9]+)\s*[+-]\s*[0-9.]+|[0-9.]+\s*\+\s*(\$?[A-Za-z]{1,3}\$?[0-9]+))\s*$"
)
# A formula using one of these changes value without the file changing, so
# its cached results are only good for the day they were computed.
_VOLATILE_RE = re.compile(r"\b(TODAY|NOW|RAND|RANDBETWEEN)\s*\(", re.IGNORECASE)


# ---------------------------------------------------------------------------
# snapshot
# ---------------------------------------------------------------------------

class _Sheet:
    __slots__ = ("name", "grid", "formulas", "date_cells", "columns", "unevaluated", "_table")

    def __init__(self, name):
        self.name = name
        self.grid = []        # every row, header included; grid[r][c] = sheet cell (r+1, c+1)
        self.formulas = {}    # (r, c) -> formula text, e.g. "=A2*50"
        self.date_cells = set()  # formula cells whose own number format is a date
        self.columns = []
        self.unevaluated = 0  # formula cells with no cached value that could not be computed
        self._table = None    # typed pyarrow table of the data rows, built on first SQL use

    @property
    def total_rows(self):
        return max(len(self.grid) - 1, 0)

    def value(self, r, c):
        row = self.grid[r] if r < len(self.grid) else ()
        return row[c] if c < len(row) else None

    def table(self):
        if self._table is None:
            self._table = _arrow_table(self)
        return self._table


class _Book:
    __slots__ = ("sheets",)

    def __init__(self, sheets):
        self.sheets = sheets  # OrderedDict name -> _Sheet, in workbook order

    def sheet(self, name=""):
        if not self.sheets:
            raise ValueError("the workbook has no sheets")
        if not name:
            return next(iter(self.sheets.values()))
        if name in self.sheets:
            return self.sheets[name]
        raise ValueError(f"no sheet named {name!r}; sheets are {list(self.sheets)}")


class TooLarge(ValueError):
    pass


def _version(file):
    st = os.stat(file)
    return (os.path.realpath(file), st.st_mtime_ns, st.st_size)


def _book(file) -> _Book:
    key = _version(file)
    if key[2] > _SNAPSHOT_MAX_BYTES:
        raise TooLarge(
            f"{os.path.basename(file)} is {key[2] / 2**20:.0f} MB; these tools load "
            f"workbooks up to {_SNAPSHOT_MAX_BYTES / 2**20:.0f} MB")
    return _build(file, key)


def evaluate(file, cache=True, sheets=None):
    """{sheet: {(r, c): value}} for every formula cell (0-based r, c): Excel's
    cached value, else pycel's. A cell neither can supply is absent.

    `sheets` limits which sheets are read at all — the editor passes only
    the small ones of a workbook whose big sheets it pages through Parquet,
    so those are never loaded into memory here. `cache` False for a scratch
    file (the editor's unsaved state), whose version is never seen again."""
    book = _build(file, _version(file), cache, None if sheets is None else frozenset(sheets))
    return {s.name: {rc: s.grid[rc[0]][rc[1]] for rc in s.formulas
                     if s.grid[rc[0]][rc[1]] is not None}
            for s in book.sheets.values()}


def formula_text(value):
    """The formula a cell holds, or None. openpyxl gives a plain formula as its
    "=..." string, but an array (CSE) formula as an ArrayFormula object and a
    what-if data table as a DataTableFormula — formulas all the same, whose
    cached result Excel saved like any other's."""
    from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

    if isinstance(value, str):
        return value if value.startswith("=") else None
    if isinstance(value, ArrayFormula):
        text = value.text or ""
        return text if text.startswith("=") else "=" + text
    if isinstance(value, DataTableFormula):
        # No formula text of its own: Excel shows it as {=TABLE(row, col)}.
        return f"=TABLE({value.r1 or ''},{value.r2 or ''})"
    return None


def _build(file, key, cache=True, only=None) -> _Book:
    sheets = OrderedDict()
    # Pass 1, data_only=False: constants, plus every formula's text and number
    # format. A workbook without formulas is complete after this pass.
    wb = openpyxl.load_workbook(file, read_only=True, data_only=False)
    try:
        epoch = wb.epoch
        for ws in wb.worksheets:
            if only is not None and ws.title not in only:
                continue
            sheet = sheets[ws.title] = _Sheet(ws.title)
            for r, cells in enumerate(ws.iter_rows()):
                row = []
                for c, cell in enumerate(cells):
                    v = cell.value
                    formula = formula_text(v)
                    if formula is not None:
                        sheet.formulas[(r, c)] = formula
                        if is_date_format(cell.number_format):
                            sheet.date_cells.add((r, c))
                        v = None
                    elif v is not None and not isinstance(v, str) and not _is_scalar(v):
                        v = str(v)
                    row.append(v)
                sheet.grid.append(row)
    finally:
        wb.close()

    volatile = any(_VOLATILE_RE.search(f) for s in sheets.values() for f in s.formulas.values())
    pending = []
    if any(s.formulas for s in sheets.values()):
        # Pass 2, data_only=True: Excel's cached result for each formula cell.
        wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                sheet = sheets.get(ws.title)
                if sheet is None or not sheet.formulas:
                    continue
                for r, row in enumerate(ws.iter_rows(values_only=True)):
                    for c, v in enumerate(row):
                        if (r, c) in sheet.formulas and v is not None:
                            sheet.grid[r][c] = v
        finally:
            wb.close()
        pending = [(s.name, r, c) for s in sheets.values()
                   for (r, c) in s.formulas if s.grid[r][c] is None]

    if pending:
        computed = _formula_values(file, key, pending, volatile, cache, only)
        for name, r, c in pending:
            sheet = sheets[name]
            if (name, r, c) not in computed:
                sheet.unevaluated += 1
                continue
            v = computed[(name, r, c)]
            if (isinstance(v, (int, float)) and not isinstance(v, bool)
                    and _shows_as_date(sheet, r, c)):
                try:
                    v = from_excel(v, epoch)
                except (ValueError, OverflowError):
                    pass  # a serial outside Excel's calendar is shown as the number it is
            sheet.grid[r][c] = v

    for sheet in sheets.values():
        sheet.columns = _column_names(sheet.grid[0] if sheet.grid else [],
                                      max((len(row) for row in sheet.grid), default=0))
    return _Book(sheets)


def _is_scalar(v):
    return isinstance(v, (bool, int, float, decimal.Decimal,
                          datetime.datetime, datetime.date, datetime.time, datetime.timedelta))


def _column_names(header, width):
    """Header names, unique case-insensitively (SQL identifiers are)."""
    names, seen = [], set()
    for j in range(width):
        v = header[j] if j < len(header) else None
        base = _text(v).strip() if v is not None else ""
        base = base or f"col{j}"
        name, n = base, 1
        while name.lower() in seen:
            n += 1
            name = f"{base}_{n}"
        seen.add(name.lower())
        names.append(name)
    return names


# ---------------------------------------------------------------------------
# formula evaluation
# ---------------------------------------------------------------------------

def _version_prefix(key):
    return f"{key[1]}-{key[2]}"


def _formula_cache_path(key, volatile, only):
    """One file per (file version, sheet subset, day for volatile books)."""
    stem = hashlib.sha1(key[0].encode("utf-8", "surrogatepass")).hexdigest()[:20]
    day = f"-{datetime.date.today().isoformat()}" if volatile else ""
    subset = "" if only is None else "-" + hashlib.sha1(
        "\0".join(sorted(only)).encode("utf-8", "surrogatepass")).hexdigest()[:12]
    return os.path.join(os.path.expanduser(_CACHE_DIR), stem,
                        f"{_version_prefix(key)}{day}{subset}.json")


def _formula_values(file, key, cells, volatile, cache=True, only=None):
    """{(sheet, r, c): value} for the uncached formula cells, from the disk
    cache or one bounded worker run. A cell the worker could not evaluate is
    absent. A run that timed out or crashed is NOT cached, so a transiently
    slow machine gets another chance next time the workbook is opened."""
    path = _formula_cache_path(key, volatile, only) if cache else None
    if path:
        try:
            with open(path, encoding="utf-8") as fh:
                return {(s, r, c): v for s, r, c, v in json.load(fh)}
        except (OSError, ValueError):
            pass
    request = json.dumps({"file": file, "cells": [[s, r + 1, c + 1] for s, r, c in cells]})
    try:
        proc = subprocess.run(
            [sys.executable, _PYCEL_WORKER],
            input=request,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=_PYCEL_TIMEOUT,
            close_fds=False,  # see executor.py's own subprocess call: avoids a
            # fork() + pthread_atfork crash when pyproj/PROJ is resident
            creationflags=subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0,
        )
    except (subprocess.TimeoutExpired, OSError):
        return {}
    if proc.returncode != 0:
        return {}
    try:
        results = json.loads(proc.stdout)
    except ValueError:
        return {}
    out = {(s, row - 1, col - 1): v for s, row, col, v in results}
    if path:
        _write_cache(path, _version_prefix(key), [[s, r, c, v] for (s, r, c), v in out.items()])
    return out


def _write_cache(path, version, rows):
    folder = os.path.dirname(path)
    try:
        os.makedirs(folder, exist_ok=True)
        for old in os.listdir(folder):  # an older version of the file is never read again
            if not old.startswith(version + "-") and not old.startswith(version + "."):
                os.remove(os.path.join(folder, old))
        tmp = f"{path}.{os.getpid()}.tmp"
        with open(tmp, "w", encoding="utf-8") as fh:
            json.dump(rows, fh, allow_nan=False)
        os.replace(tmp, path)
    except OSError:
        pass  # a cache that can't be written costs speed, not correctness


def _shows_as_date(sheet, r, c, depth=0):
    """Whether Excel would display the computed number at (r, c) as a date:
    a date-formatted or whole-date-call formula, a date constant — or an
    offset off one of those (one level: `due = issued + 30`)."""
    formula = sheet.formulas.get((r, c))
    if formula is None:
        return isinstance(sheet.value(r, c), (datetime.date, datetime.datetime))
    if (r, c) in sheet.date_cells or _is_date_call(formula):
        return True
    m = _OFFSET_RE.match(formula.strip())
    if not m or depth:
        return False
    row, col = coordinate_to_tuple((m.group(1) or m.group(2)).replace("$", "").upper())
    return _shows_as_date(sheet, row - 1, col - 1, depth + 1)


def _is_date_call(formula):
    """True only when the WHOLE formula is one date-function call —
    `TODAY()-A2` starts like one but is a day count."""
    body = formula.lstrip("=").strip()
    upper = body.upper()
    for name in _DATE_FUNCS:
        if not upper.startswith(name):
            continue
        depth = 0
        for i in range(len(name) - 1, len(body)):
            if body[i] == "(":
                depth += 1
            elif body[i] == ")":
                depth -= 1
                if depth == 0:
                    return body[i + 1:].strip() == ""
        return False
    return False


# ---------------------------------------------------------------------------
# values out
# ---------------------------------------------------------------------------

def _jsonify(value):
    """A cell value as JSON: dates ISO, Decimal/bytes/timedelta as text."""
    if value is None or isinstance(value, (bool, int, str)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else "#NUM!"
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, (bytes, bytearray)):
        return value.hex()
    return str(value)


def _text(value):
    v = _jsonify(value)
    return v if isinstance(v, str) else json.dumps(v)


def _column_kind(values):
    """number / integer / date / boolean / text / empty — what SQL sees.
    Error values (#DIV/0!, …) and blank strings don't decide a column's type;
    they read as NULL in a typed one, the way Excel's SUM skips them."""
    kinds = set()
    for v in values:
        if v is None or (isinstance(v, str) and (v in _EXCEL_ERRORS or not v.strip())):
            continue
        if isinstance(v, bool):
            kinds.add("boolean")
        elif isinstance(v, int):
            kinds.add("integer")
        elif isinstance(v, (float, decimal.Decimal)):
            kinds.add("number")
        elif isinstance(v, (datetime.datetime, datetime.date)):
            kinds.add("date")
        else:
            kinds.add("text")
        if len(kinds) > 1 and kinds != {"integer", "number"}:
            return "text"
    if not kinds:
        return "empty"
    return "number" if kinds == {"integer", "number"} else kinds.pop()


def _typed_column(values, kind):
    import pyarrow as pa

    if kind == "text":
        return pa.array([None if v is None else _text(v) for v in values], pa.string())
    vals = [None if v is None or (isinstance(v, str) and (v in _EXCEL_ERRORS or not v.strip()))
            else v for v in values]
    if kind == "integer" and all(v is None or -2**63 <= v < 2**63 for v in vals):
        return pa.array(vals, pa.int64())
    if kind in ("integer", "number"):
        return pa.array([None if v is None else float(v) for v in vals], pa.float64())
    if kind == "boolean":
        return pa.array(vals, pa.bool_())
    if kind == "date":
        return pa.array([v if v is None or isinstance(v, datetime.datetime)
                         else datetime.datetime(v.year, v.month, v.day) for v in vals],
                        pa.timestamp("us"))
    return pa.nulls(len(values), pa.string())  # empty


def _arrow_table(sheet):
    import pyarrow as pa

    data = sheet.grid[1:]
    arrays = []
    for c in range(len(sheet.columns)):
        values = [row[c] if c < len(row) else None for row in data]
        arrays.append(_typed_column(values, _column_kind(values)))
    return pa.table(arrays, names=sheet.columns)


def _schema(sheet):
    table = sheet.table()
    return [{"name": f.name, "type": str(f.type)} for f in table.schema]


# ---------------------------------------------------------------------------
# tools
# ---------------------------------------------------------------------------

def sheets(file: str) -> dict:
    """List every sheet in the workbook with its row count and typed columns.
    Start here: the sheet names are the table names `query` accepts."""
    book = _book(file)
    return {"file": file, "sheets": [
        {"name": s.name, "rows": s.total_rows, "columns": _schema(s),
         "formulas": len(s.formulas), "unevaluated": s.unevaluated}
        for s in book.sheets.values()]}


def describe(file: str, sheet: str = "", columns: str = "") -> dict:
    """Summary statistics for each column of a sheet (the first by default):
    count, nulls, distinct; min/max/mean/median/std/sum for numbers; min/max
    for dates; the most common values for text. `columns` is an optional
    comma-separated subset."""
    import duckdb

    book = _book(file)
    s = book.sheet(sheet)
    table = s.table()
    wanted = [c.strip() for c in columns.split(",") if c.strip()] if columns else table.column_names
    missing = [c for c in wanted if c not in table.column_names]
    if missing:
        raise ValueError(f"no column(s) {missing} in sheet {s.name!r}; columns are {table.column_names}")

    con = _sandbox(duckdb, {"t": table})
    try:
        stats = []
        for name in wanted:
            q = _ident(name)
            kind = str(table.schema.field(name).type)
            aggs = [f"count({q})", f"count(*) - count({q})", f"count(DISTINCT {q})"]
            keys = ["count", "nulls", "distinct"]
            if kind in ("int64", "double"):
                aggs += [f"min({q})", f"max({q})", f"avg({q})", f"median({q})",
                         f"stddev_samp({q})", f"sum({q})"]
                keys += ["min", "max", "mean", "median", "std", "sum"]
            elif kind.startswith("timestamp"):
                aggs += [f"min({q})", f"max({q})"]
                keys += ["min", "max"]
            row = con.execute(f"SELECT {', '.join(aggs)} FROM t").fetchone()
            entry = {"column": name, "type": kind}
            entry.update({k: _jsonify(v) for k, v in zip(keys, row)})
            if kind == "string":
                entry["top"] = [
                    {"value": v, "count": n} for v, n in con.execute(
                        f"SELECT {q}, count(*) AS n FROM t WHERE {q} IS NOT NULL "
                        f"GROUP BY 1 ORDER BY n DESC, 1 LIMIT {_TOP_VALUES}").fetchall()]
            c = s.columns.index(name)
            errors = sum(1 for r in range(1, len(s.grid))
                         if isinstance(v := s.value(r, c), str) and v in _EXCEL_ERRORS)
            if errors:
                entry["errors"] = errors
            stats.append(entry)
    finally:
        con.close()
    return {"sheet": s.name, "rows": s.total_rows, "unevaluated": s.unevaluated, "columns": stats}


def query(file: str, sql: str, limit: int = 1000) -> dict:
    """Run read-only DuckDB SQL over the workbook. Every sheet is a table
    named after it (quote names with spaces: SELECT * FROM "Q1 sales");
    column names come from each sheet's header row — `sheets` lists them with
    their types. At most `limit` rows are returned; `truncated` says whether
    there were more."""
    import duckdb

    book = _book(file)
    limit = min(max(int(limit), 1), _QUERY_MAX_ROWS)
    con = _sandbox(duckdb, {s.name: s.table() for s in book.sheets.values()})
    timer = threading.Timer(_QUERY_TIMEOUT, con.interrupt)
    timer.start()
    try:
        cur = con.execute(sql)
        if cur.description is None:
            return {"columns": [], "types": [], "rows": [], "truncated": False}
        cols = [d[0] for d in cur.description]
        types = [str(d[1]) for d in cur.description]
        fetched = cur.fetchmany(limit + 1)
    except duckdb.InterruptException:
        raise TimeoutError(f"query exceeded {_QUERY_TIMEOUT:g}s") from None
    finally:
        timer.cancel()
        con.close()
    return {
        "columns": cols,
        "types": types,
        "rows": [[_jsonify(v) for v in row] for row in fetched[:limit]],
        "truncated": len(fetched) > limit,
    }


def cells(file: str, ref: str, sheet: str = "") -> dict:
    """The cells of an A1 reference ("B5", "A1:D10", or "Sheet 2!A1:C3") as
    rows of {ref, value, formula}. Rows are the sheet's own numbering, header
    included — the way the formulas themselves address cells."""
    book = _book(file)
    if "!" in ref:
        sheet, ref = ref.rsplit("!", 1)
        sheet = sheet.strip("'").replace("''", "'")
    s = book.sheet(sheet)
    min_col, min_row, max_col, max_row = range_boundaries(ref.replace("$", "").upper())
    if min_row is None or min_col is None:
        raise ValueError(f"{ref!r} is not a bounded A1 range")
    if (max_row - min_row + 1) * (max_col - min_col + 1) > _CELLS_MAX:
        raise ValueError(f"{ref!r} spans more than {_CELLS_MAX} cells; use query for bulk reads")
    rows = []
    for row in range(min_row, max_row + 1):
        rows.append([
            {"ref": f"{get_column_letter(col)}{row}",
             "value": _jsonify(s.value(row - 1, col - 1)),
             "formula": s.formulas.get((row - 1, col - 1))}
            for col in range(min_col, max_col + 1)])
    return {"sheet": s.name, "ref": ref, "cells": rows}


def _ident(name):
    return '"' + name.replace('"', '""') + '"'


def _sandbox(duckdb, tables):
    """An in-memory DuckDB that sees `tables` and nothing else: no files, no
    network, no extension installs, bounded memory and threads — then locked,
    so the SQL it runs can't turn any of that back on."""
    con = duckdb.connect(":memory:")
    try:
        for name, table in tables.items():
            con.register(name, table)
        con.execute("SET memory_limit='1GB'")
        con.execute("SET threads=2")
        con.execute("SET enable_external_access=false")
        con.execute("SET autoinstall_known_extensions=false")
        con.execute("SET autoload_known_extensions=false")
        con.execute("SET lock_configuration=true")
    except BaseException:
        con.close()
        raise
    return con
